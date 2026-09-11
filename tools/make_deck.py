# make_deck.py - 15-slide academic presentation (W29, ADR-021).
# Assembles existing figures + ADR lineage into presentation.pptx.
# Slide count/notes are speaker support; numbers pulled from git.
import subprocess
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
FIGS = ROOT / "docs" / "figures"
SHOTS = ROOT / "data" / "logs"

def _git_count(args):
    r = subprocess.run(["git"] + args, cwd=ROOT,
                       capture_output=True, text=True)
    return len([l for l in r.stdout.splitlines() if l.strip()])

N_COMMITS = _git_count(["log", "--oneline", "--all"])
N_TAGS = _git_count(["tag", "--list"])

prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)


def slide(title, bullets, image=None, notes=""):
    s = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    tb = s.shapes.add_textbox(Inches(0.4), Inches(0.3),
                              Inches(12.5), Inches(1.0))
    p = tb.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(30)
    p.font.bold = True
    body = s.shapes.add_textbox(Inches(0.4), Inches(1.5),
                                Inches(6.2) if image else Inches(12.5),
                                Inches(5.4))
    tf = body.text_frame
    tf.word_wrap = True
    for i, b in enumerate(bullets):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.text = chr(8226) + " " + b
        para.font.size = Pt(18)
    if image and image.exists():
        s.shapes.add_picture(str(image), Inches(6.8), Inches(1.5),
                             width=Inches(6.2))
    s.notes_slide.notes_text_frame.text = notes
    return s

# ---- Slide 1: title ----
slide("Hybrid-Brain: A Test-Driven Spiking Cognitive Architecture",
      ["Integrating reward learning, affect, personality and language",
       f"{N_COMMITS} commits | {N_TAGS} tagged milestones | "
       "111 tests | 21 architecture decision records",
       "Every number in this deck regenerates from a command"],
      notes="Open with the methodology: test-driven, every acceptance "
            "criterion measured, 21 ADRs. The claims are computational, "
            "not phenomenal - stated up front.")

# ---- Slide 2: the integration gap ----
slide("The integration gap",
      ["Spiking simulators (Brian2, NEST): rigorous neural dynamics,",
       "   but no cognition, affect or behavior loop",
       "Cognitive agents (ACT-R, Generative Agents): rich cognition,",
       "   but no neural substrate and no reward learning",
       "Hybrid-Brain: connects spiking cortex -> affect -> personality",
       "   -> behavior in ONE loop, every module traced to literature"],
      notes="Name the two families and position between them. The "
            "mapping table (next slide) is the evidence of the trace.")

# ---- Slide 3: module -> region -> literature ----
slide("Every module traced to cognitive neuroscience",
      ["Dopaminergic RPE -> Schultz, Dayan & Montague (1997) Science",
       "Place cells -> O'Keefe & Nadel (1978) Cognitive Map",
       "Somatic markers -> Damasio (1996) Phil Trans R Soc B",
       "Selective attention -> Posner & Petersen (1990) Annu Rev",
       "Habit vs goal -> Balleine & Dickinson (1998) Neuropharm",
       "Full table: docs/neuroscience_mapping.md (20+ modules)"],
      notes="Point to the repo file live if needed. Supporting refs "
            "marked honestly in the table.")

# ---- Slide 4: spiking core + calibration story ----
slide("Methods honesty: measurement over derivation",
      ["LIF neurons + E/I synapses, validated against Brian2",
       "f-I threshold: computed 0.93 (theory) vs measured 10-20",
       "The measurement superseded the derivation - recorded as a",
       "   methodological rule: data over derivation",
       "Amplifies to every module: acceptance by number, not claim"],
      notes="This slide pre-empts the 'is it rigorous' question. "
            "The calibration test that settled a two-theory dispute "
            "is in tests/ and git history.")

# ---- Slide 5: regions ----
slide("Implemented brain regions",
      ["Thalamus: deterministic sensory relay",
       "Hippocampus: place cells (2D Gaussian fields)",
       "Amygdala: threat conditioning, dopamine-gated 3-factor rule",
       "Cortex6: six layers, homeostatic row-sum feedforward",
       "Basal ganglia: Go/No-Go channels, cross-inhibition",
       "PFC: recurrent working memory with delay persistence"],
      notes="Each region passed sweep or acceptance tests; cortex6 "
            "was tuned via a 16-combo sweep (tune_cortex.py).")

# ---- Slide 6: emotion dynamics ----
slide("Emotion dynamics under sustained failure",
      ["Six Ekman emotions with timescale hierarchy:",
       "   dopamine ~100ms | emotions 2-30s | moods ~hours (ADR-012)",
       "sadness (tau 30s) rises slowly and persists; anger faster",
       "   (blocked-escape profile) - matches figure"],
      image=FIGS / "fig1_emotion_dynamics.png",
      notes="The hierarchy is TESTABLE (test_timescale_hierarchy) "
            "and was the tripwire that caught a real regression.")

# ---- Slide 7: live emotion monitor ----
slide("Live affect monitor + narration",
      ["struggle scenario: sadness 0->0.16, anger 0->0.30",
       "valence -0.36, arousal 0.86 (measured run)",
       "Narration reflects the dominant emotion - telemetry in",
       "   language, per ADR-018"],
      image=SHOTS / "struggle.PNG",
      notes="The screenshot is a real run. If asked: the emotion "
            "levels feed somatic markers which bias action values.")

# ---- Slide 8: cocktail party ----
slide("Selective attention with two real streams",
      ["Visual cortex vs auditory stream, genuinely competing",
       "Attended stream gets 3x gain; dwell prevents flicker",
       "Loud auditory event STEALS attention after dwell",
       "Posner & Petersen (1990) networks, implemented"],
      image=FIGS / "fig2_cocktail_party.png",
      notes="The steal is deterministic and test-verified "
            "(test_loud_event_steals_attention).")

# ---- Slide 9: language boundary ----
slide("Language narrates, never decides (ADR-018)",
      ["An LLM that acts = an untrained foreign brain: outside RPE,",
       "   plasticity, consolidation - nothing teaches it",
       "Language = READ-ONLY telemetry: consumes GlobalWorkspace,",
       "   emotions, attention; emits first-person narration",
       "First spoken sentence from real brain state"],
      image=SHOTS / "brain spoke for the first time..PNG",
      notes="This boundary protects the attribution chain - "
            "the 'why this action' question stays answerable.")

# ---- Slide 10: habit dissociation ----
slide("Habit vs goal-directed dissociation (ADR-019)",
      ["Habits form only on LOW RPE-variance (stable values)",
       "Ballistic execution: no value lookup, no BG call",
       "Under reward devaluation: inertia persists, then an",
       "   accumulated-error detector invalidates the cache",
       "Balleine & Dickinson (1998); Yin & Knowlton (2006)"],
      image=FIGS / "fig3_habit_dissociation.png",
      notes="The dissociation is the point: two control regimes, "
            "measurably different. Dossier Q7 extends this.")

# ---- Slide 11: somatic markers ----
slide("Somatic markers bias decisions (Damasio 1996)",
      ["Punishment + arousal accrue a bounded position penalty",
       "Read-time bias on action utilities - never written into",
       "   the value table (extinction snaps to zero past threshold)",
       "Ablation-tested: avoidance with ZERO navigation cost"],
      image=FIGS / "fig4_somatic_marker.png",
      notes="Week 18 acceptance: ON-arm 0/10 vs OFF-arm 10/10 visits "
            "to the punished region, eval goals 10/10 both arms.")

# ---- Slide 12: navigation ----
slide("Visual-motor closed loop",
      ["VisualWorld render -> MobileNet features -> six-layer cortex",
       "-> attention -> Q(cell, action) -> MotorActuator -> reward",
       "Best evaluation trajectory: distance 1.50 -> 0.50",
       "Honest scope: decision via proven Q-table; neural readout",
       "   live in-loop and logged"],
      image=FIGS / "fig5_navigation.png",
      notes="The scope note is deliberate - attribution clarity over "
            "claim inflation. Full neural readout is future work.")

# ---- Slide 13: EEG ----
slide("Simulated EEG: rest vs threat",
      ["Region activity -> 10-20 montage channels -> FFT band power",
       "Threat shifts BETA power x1.94 (stress-EEG consistent)",
       "LIMITATION stated: rest gamma inflated by constant 75Hz",
       "   drive; simulated EEG proxy, not a scalp forward model"],
      image=FIGS / "fig6_eeg_spectrum.png",
      notes="State the limitation BEFORE being asked. The gamma "
            "refinement (rhythmic rest drive) is future work.")

# ---- Slide 14: limitations ----
slide("Honest limitations",
      ["Affect is COMPUTATIONAL, not phenomenal - no qualia claim",
       "Rat-scale abstractions: grid world, tone events",
       "Single-seed measurements; test-retest pending",
       "EEG is a simulated proxy; gamma baseline limitation",
       "Scope decisions recorded: ADR-006, ADR-020, ADR-021",
       "   (superhuman features deferred, not silently cut)"],
      notes="This slide builds credibility with academics - the "
            "limitation list is as important as the claims.")

# ---- Slide 15: reproducibility ----
slide("Reproducibility",
      [f"{N_COMMITS} commits, {N_TAGS} tags, 111 tests, 21 ADRs",
       "Seed registry: every RNG seed scanned and documented",
       "Figure regeneration: one command each (make_figures,",
       "   eeg_simulator, demo_month6)",
       "Repository: github.com/esmaeilireza/Hybrid-Brain"],
      notes="Close: every number shown today regenerates from a "
            "command. Invite the committee to run the suite.")

out = ROOT / "docs" / "paper" / "presentation.pptx"
prs.save(str(out))
print("DECK SAVED:", out, "-", len(prs.slides.__iter__.__self__._sldIdLst), "slides")

